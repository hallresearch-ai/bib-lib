"""Apply reviewed per-name author decisions and synchronize auditable exports."""
import csv,json,shutil
from collections import Counter
from copy import copy
from pathlib import Path
import openpyxl,bibtexparser
from openpyxl.cell.rich_text import CellRichText,TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill,Font
from resolve_authors import parts,norm
BASE=Path('bib_files/normalized-2026-10-08')
GOOD={'confirmed','corrected','normalized_by_user'}
def main():
 path=BASE/'library-cleaned.xlsx';shutil.copy2(path,'/tmp/library-cleaned-before-author-verification.xlsx')
 before=openpyxl.load_workbook(path,rich_text=True);w=openpyxl.load_workbook(path,rich_text=True);s=w['Bibliography'];headers=[c.value for c in s[1]];ki=headers.index('citation_key');ai=headers.index('author')
 decisions=json.loads((BASE/'author-verification-decisions.json').read_text());assert len(decisions)==637
 review=w['Review'];rh=[c.value for c in review[1]];audit=w['Change audit'];resolved=0
 # Resolve only the author-name variants supported by the publication-specific decisions.
 for r in review.iter_rows(min_row=2):
  if r[3].value!='author' or r[5].value!='open' or r[2].value!='author_variant':continue
  key=r[0].value;detail=r[4].value or '';variants=[parts(v) for v in detail.split('; ')]
  def relevant(n):
   for text in [n.get('before',''),n['after']]:
    if not text or text=='others':continue
    p=parts(text)
    for v in variants:
     if norm(p['family'])==norm(v['family']) and (not p['given'] or not v['given'] or norm(p['given'])[:1]==norm(v['given'])[:1]):return True
   return False
  names=[n for n in decisions[key]['names'] if relevant(n)]
  if names and all(n['status'] in GOOD for n in names):
   evidence='; '.join(dict.fromkeys(n['source_url'] for n in names if n['source_url']))
   message='Publication-specific evidence confirms '+', '.join(dict.fromkeys(n['after'] for n in names))+'. Same-surname authors are not merged; supported initials and spelling corrections are applied.'
   r[5].value='resolved';r[6].value=message;r[7].value=evidence;resolved+=1
   audit.append([key,'author_verification','open','resolved',message,evidence])
 av=w.create_sheet('Author verification');av.append(['citation_key','previous_name','normalized_name','status','verification_method','source_url','external_name','orcid','decision_reason'])
 for c in av[1]:c.font=Font(bold=True,color='FFFFFFFF');c.fill=PatternFill('solid',fgColor='243E50')
 av.freeze_panes='A2';changed=0;black_names=0;red_names=0
 for row in s.iter_rows(min_row=2):
  key=row[ki].value;c=row[ai];d=decisions[key];old=str(c.value) if c.value is not None else None;assert old==d['before'],key
  if d['after']!=old:
   audit.append([key,'author',old,d['after'],'Author list corrected from publication-specific sources; before/after names and source evidence in Author verification','; '.join(dict.fromkeys(n['source_url'] for n in d['names'] if n['source_url']))]);changed+=1
  uncertain=[n for n in d['names'] if n['status'] not in GOOD]
  if uncertain or not d['names']:
   detail='Unverified names: '+('; '.join(n['after'] for n in uncertain) if uncertain else '(empty author field)')
   detail+='.'
   evidence='; '.join(dict.fromkeys(n['source_url'] for n in uncertain if n['source_url']))
   review.append([key,'accuracy','author_verification_pending','author',detail,'open',None,evidence or None])
   c.comment=Comment('Remaining author verification findings:\n\n'+'\n\n'.join(n['after']+': '+n['reason'] for n in uncertain) if uncertain else 'Author field is empty; authorship could not be verified.','Audit')
  else:c.comment=None
  font=copy(c.font);font.color='FF000000';c.font=font
  if not d['names']:
   c.value=None;c.fill=PatternFill('solid',fgColor='FFFFC7CE');av.append([key,None,None,'unverified','No author supplied',None,None,None,'Author field is empty; no supported author list found.']);continue
  runs=[]
  for i,n in enumerate(d['names']):
   if i:runs.append(' and ')
   color='FF000000' if n['status'] in GOOD else 'FFFF0000'
   runs.append(TextBlock(InlineFont(color=color),n['after']))
   av.append([key,n.get('before'),n['after'],n['status'],n['method'],n['source_url'],n.get('external_name'),n.get('orcid'),n['reason']])
   if color=='FF000000':black_names+=1
   else:red_names+=1
  if all(n['status'] in GOOD for n in d['names']):c.value=d['after']
  elif all(n['status'] not in GOOD for n in d['names']):c.value=d['after'];font=copy(c.font);font.color='FFFF0000';c.font=font
  else:c.value=CellRichText(runs)
  if c.fill.patternType=='solid' and c.fill.fgColor.type=='rgb' and c.fill.fgColor.rgb in {'FFFF0000','00FF0000','FFFFC7CE','00FFC7CE'}:c.fill=PatternFill()
 av.auto_filter.ref=av.dimensions
 for col,width in {'A':28,'B':33,'C':33,'D':23,'E':42,'F':75,'G':33,'H':45,'I':100}.items():av.column_dimensions[col].width=width
 # No title values, title markings, other bibliography values or other styling may change.
 for oldrow,newrow in zip(before['Bibliography'],s):
  for x,y in zip(oldrow,newrow):
   if x.row>1 and x.column==ai+1:continue
   assert str(x.value)==str(y.value) and copy(x._style)==copy(y._style) and x.comment==y.comment,x.coordinate
 rows=[dict(zip(headers,[str(c.value) if isinstance(c.value,CellRichText) else c.value for c in row])) for row in s.iter_rows(min_row=2)]
 control={'entry_type','citation_key'};bib='\n\n'.join('@'+r['entry_type']+'{'+r['citation_key']+',\n'+',\n'.join('  '+f+' = {'+str(r[f])+'}' for f in headers if f not in control and r.get(f))+'\n}' for r in rows)+'\n'
 parsed=bibtexparser.parse_string(bib);assert not parsed.failed_blocks and len(parsed.entries)==637
 actual={e.key:(e.entry_type,{f.key:f.value for f in e.fields}) for e in parsed.entries}
 for r in rows:assert actual[r['citation_key']]==(r['entry_type'],{f:str(r[f]) for f in headers if f not in control and r.get(f)}),r['citation_key']
 policy='Author verification update: every record was researched through matched DOI author lists, publisher/repository author metadata and opening-page PDF author blocks, with 180 public ORCID records read to support identity checks. Initials expand only with publication-specific evidence. Full author lists are restored when supported without changing the cited version. Same-surname people remain distinct. Confirmed/corrected names and prior user-approved Hall, Patrick normalization are black; unsupported names are red individually, and empty author cells have red fill. Each name decision and source is recorded in Author verification and author-verification-decisions.json. All non-author bibliography values and title markings are preserved.'
 w['Policy'].append([policy]);tmp=path.with_name('author-verification.tmp.xlsx');w.save(tmp)
 check=openpyxl.load_workbook(tmp,rich_text=True)
 for r in check['Bibliography'].iter_rows(min_row=2):
  key=r[ki].value;c=r[ai];d=decisions[key];assert (str(c.value) if c.value is not None else None)==d['after'],key
  if isinstance(c.value,CellRichText):
   chunks=[x for x in c.value if isinstance(x,TextBlock)];assert len(chunks)==len(d['names']),key
   for chunk,n in zip(chunks,d['names']):assert chunk.text==n['after'] and chunk.font.color.rgb==('FF000000' if n['status'] in GOOD else 'FFFF0000'),key
  elif d['names']:assert c.font.color.rgb==('FF000000' if all(n['status'] in GOOD for n in d['names']) else 'FFFF0000'),key
  else:assert c.fill.fgColor.rgb=='FFFFC7CE'
 tmp.replace(path);(BASE/'library-cleaned.bib').write_text(bib)
 for sn,fn in [('Review','review.csv'),('Change audit','change-audit.csv'),('Author verification','author-verification.csv')]:
  with (BASE/fn).open('w',newline='',encoding='utf-8') as f:csv.writer(f).writerows(w[sn].values)
 with (BASE/'normalization-policy.txt').open('a') as f:f.write('\n'+policy+'\n')
 m=json.loads((BASE/'manifest.json').read_text());rv=[dict(zip(rh,r)) for r in list(review.values)[1:]]
 m['field_changes']=audit.max_row-1;m['review_status_counts']=dict(Counter(i['status'] for i in rv));m['review_counts']=dict(Counter(i['category'] for i in rv if i['status']=='open'));m['qc_counts']=dict(Counter(i['check'] for i in rv if i['status']=='open'))
 statuses=Counter(n['status'] for r in decisions.values() for n in r['names'])
 m['author_verification']={'records_checked':637,'fully_externally_verified_author_fields':sum(bool(r['names']) and all(n['status'] in {'confirmed','corrected'} for n in r['names']) for r in decisions.values()),'changed_author_fields':changed,'full_lists_restored':sum(r['complete_author_list_restored'] for r in decisions.values()),'newly_resolved_variant_findings':resolved,'name_status_counts':dict(statuses),'black_name_occurrences':black_names,'red_name_occurrences':red_names,'empty_author_fields':sum(not r['names'] for r in decisions.values()),'orcid_public_records_read':180,'decisions_file':'author-verification-decisions.json','evidence_file':'author-verification-evidence.jsonl','name_colors_verified':True,'non_author_cells_unchanged_verified':True}
 (BASE/'manifest.json').write_text(json.dumps(m,indent=2)+'\n');print(json.dumps(m['author_verification'],indent=2))
if __name__=='__main__':main()
