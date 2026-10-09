"""Apply separately reviewed title decisions, preserving all non-title data/styles."""
import csv,json,shutil
from collections import Counter
from copy import copy
from pathlib import Path
import openpyxl,bibtexparser
from openpyxl.styles import PatternFill
BASE=Path('bib_files/normalized-2026-10-08')
def main():
 decisions=json.loads((BASE/'title-followup-decisions.json').read_text())
 path=BASE/'library-cleaned.xlsx'
 before=openpyxl.load_workbook(path)
 w=openpyxl.load_workbook(path)
 s=w['Bibliography'];headers=[c.value for c in s[1]];ci=headers.index('citation_key');ti=headers.index('title')
 keys={row[ci].value:row for row in s.iter_rows(min_row=2)}
 review=w['Review'];rh=[c.value for c in review[1]];ri={h:i for i,h in enumerate(rh)}
 pending={r[0].value for r in review.iter_rows(min_row=2) if r[ri['field']].value=='title' and r[ri['status']].value=='open'}
 assert set(decisions)<=pending, 'Decisions must refer to currently pending titles'
 shutil.copy2(path,'/tmp/library-cleaned-before-title-followup.xlsx')
 audit=w['Change audit']
 for k,d in decisions.items():
  c=keys[k][ti];old=c.value
  if d['action']=='corrected':
   c.value=d['new_title'];audit.append([k,'title',old,c.value,d['reason'],d['source_url']])
  font=copy(c.font);font.color='FF000000';c.font=font
  if c.fill.patternType=='solid' and c.fill.fgColor.type=='rgb' and c.fill.fgColor.rgb in {'FFFF0000','00FF0000','FFFFC7CE','00FFC7CE'}:c.fill=PatternFill()
  c.comment=None
  audit.append([k,'title_verification','open','resolved',d['action']+': '+d['reason'],d['source_url']])
 for r in review.iter_rows(min_row=2):
  k=r[0].value
  if k in decisions and r[ri['field']].value=='title' and r[ri['status']].value=='open':
   d=decisions[k];r[ri['status']].value='resolved';r[ri['resolution']].value=d['action'].capitalize()+': '+d['reason'];r[ri['evidence']].value=d['source_url']
 evidence=list(map(json.loads,(BASE/'title-verification.jsonl').read_text().splitlines()))
 for r in evidence:
  k=r['citation_key']
  if k in decisions:
   d=decisions[k];r['previous_workbook_title']=r['workbook_title'];r['workbook_title']=keys[k][ti].value;r['status']='matched';r['followup_action']=d['action'];r['followup_reason']=d['reason']
   r['evidence'].append({'source_url':d['source_url'],'final_url':d['source_url'],'kind':d['method'],'title':d['external_title'],'matches':True,'reviewed':True})
 v=w['Title verification']
 for r in v.iter_rows(min_row=2):
  k=r[0].value
  if k in decisions:
   d=decisions[k];r[1].value=keys[k][ti].value;r[2].value='resolved';r[3].value=d['external_title'];r[4].value=d['source_url'];r[5].value=d['method'];r[6].value=d['action'].capitalize()+': '+d['reason']
 policy='Title follow-up: all 254 pending titles were researched using PDF opening pages, full publication pages and title/author/year registry searches. Reviewed evidence confirms 143 titles and corrects 11. Capitalization, braces, punctuation, explicit report identifiers and descriptive prefixes are harmless when identity matches. Original cited versions are preserved. Resolved title cells have black text and no audit comment; 100 unsupported titles remain red. Decisions and source URLs are recorded in title-followup-decisions.json; research evidence is in title-followup-evidence.jsonl.'
 w['Policy'].append([policy])
 rows=[dict(zip(headers,[c.value for c in row])) for row in s.iter_rows(min_row=2)]
 control={'entry_type','citation_key'}
 bib='\n\n'.join('@'+r['entry_type']+'{'+r['citation_key']+',\n'+',\n'.join('  '+f+' = {'+str(r[f])+'}' for f in headers if f not in control and r.get(f))+'\n}' for r in rows)+'\n'
 parsed=bibtexparser.parse_string(bib);assert not parsed.failed_blocks;assert len(parsed.entries)==637
 actual={e.key:(e.entry_type,{f.key:f.value for f in e.fields}) for e in parsed.entries}
 for r in rows:assert actual[r['citation_key']]==(r['entry_type'],{f:str(r[f]) for f in headers if f not in control and r.get(f)}),r['citation_key']
 # Verify every unaffected bibliography cell, including existing audit styling.
 for oldrow,newrow in zip(before['Bibliography'],s):
  for a,b in zip(oldrow,newrow):
   if a.row>1 and a.column==ti+1 and oldrow[ci].value in decisions:continue
   assert a.value==b.value and copy(a._style)==copy(b._style) and a.comment==b.comment,(a.coordinate,'unexpected change')
 tmp=path.with_name('library-cleaned.title-followup.tmp.xlsx');w.save(tmp)
 loaded=openpyxl.load_workbook(tmp)
 for r in loaded['Bibliography'].iter_rows(min_row=2):
  if r[ci].value in decisions:assert r[ti].font.color.rgb=='FF000000' and r[ti].comment is None
 tmp.replace(path)
 (BASE/'library-cleaned.bib').write_text(bib)
 (BASE/'title-verification.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in evidence))
 for sheet,filename in [('Review','review.csv'),('Change audit','change-audit.csv'),('Title verification','title-verification.csv')]:
  with (BASE/filename).open('w',newline='',encoding='utf-8') as f:csv.writer(f).writerows(w[sheet].values)
 with (BASE/'normalization-policy.txt').open('a') as f:f.write('\n'+policy+'\n')
 m=json.loads((BASE/'manifest.json').read_text());reviews=[dict(zip(rh,r)) for r in review.values if r[0]!='citation_key']
 m['field_changes']=audit.max_row-1;m['review_status_counts']=dict(Counter(r['status'] for r in reviews))
 m['review_counts']=dict(Counter(r['category'] for r in reviews if r['status']=='open'))
 m['qc_counts']=dict(Counter(r['check'] for r in reviews if r['status']=='open'))
 m['title_verification'].update(resolved=466,pending_wording_or_version_verification=100,followup_confirmed=143,followup_corrected=11,followup_decisions_file='title-followup-decisions.json',followup_evidence_file='title-followup-evidence.jsonl')
 m['title_followup_non_title_cells_unchanged_verified']=True;m['roundtrip_verified']=True
 (BASE/'manifest.json').write_text(json.dumps(m,indent=2)+'\n')
 print(json.dumps(m['title_verification'],indent=2))
if __name__=='__main__':main()
