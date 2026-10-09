"""Apply URL checks without treating mutability as a bibliography error."""
import csv,json,shutil
from collections import Counter
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font,PatternFill
B=Path('bib_files/normalized-2026-10-08')
def main():
 path=B/'library-cleaned.xlsx';shutil.copy2(path,'/tmp/library-cleaned-before-url-verification.xlsx');old=openpyxl.load_workbook(path,rich_text=True);w=openpyxl.load_workbook(path,rich_text=True);s=w['Bibliography'];h=[c.value for c in s[1]];ui=h.index('url');ki=h.index('citation_key');rows={r[ki].value:r for r in s.iter_rows(min_row=2)}
 evidence={r['citation_key']:r for r in map(json.loads,(B/'url-verification-evidence.jsonl').read_text().splitlines())};assert len(evidence)==637
 review=w['Review'];rh=[c.value for c in review[1]];original=[dict(zip(rh,r)) for r in list(review.values)[1:]];flagged={r['citation_key'] for r in original if r['field']=='url' and r['status']=='open'};audit=w['Change audit'];removed=0;cleared=set()
 for r in review.iter_rows(min_row=2):
  if r[3].value!='url' or r[5].value!='open':continue
  key=r[0].value;ev=evidence[key]
  if r[2].value=='mutable_web_resource':
   r[5].value='resolved';r[6].value='Mutability alone is not an issue under the user-approved URL policy. Actual resolution and record-match results are tracked separately.';r[7].value='Explicit user instruction; URL verification';removed+=1
   audit.append([key,'url_verification','mutable_web_resource','resolved','Remove mutability-only warning; record actual URL check separately',r[7].value])
  elif ev['status']=='matched':
   r[5].value='resolved';r[6].value='URL resolves to content matching the cited resource; generic homepages are accepted when the homepage is the cited resource.';r[7].value=ev.get('final_url');audit.append([key,'url_verification','open','resolved',r[6].value,r[7].value])
 # Only definite new errors are added automatically. An automated access block alone
 # is recorded as information unless an existing URL flag is awaiting verification.
 known_wrong={'DiscussionofBoosting173','GenerativeRedTeamCha266','SupersparseLinearInt536','FacebooksFivePillars224'}
 counts=Counter()
 for key,row in rows.items():
  ev=evidence[key];c=row[ui];status=ev['status'];assert c.value==ev['requested_url'],key
  matched=status=='matched';needs_flag=not matched and (key in flagged or status in {'http_error','missing'} or key in known_wrong)
  if matched:
   font=copy(c.font);font.color='FF000000';c.font=font;c.comment=None
   if c.fill.patternType=='solid' and c.fill.fgColor.type=='rgb' and c.fill.fgColor.rgb in {'FFFF0000','00FF0000','FFFFC7CE','00FFC7CE'}:c.fill=PatternFill()
   if key in flagged:cleared.add(key)
  elif needs_flag:
   check='url_http_error' if status=='http_error' else 'missing_url' if status=='missing' else 'url_record_mismatch' if key in known_wrong else 'url_verification_pending'
   if status=='http_error':detail='URL returned HTTP '+str(ev.get('http_status',''))+'. '+ '; '.join(ev['notes'])
   elif status=='missing':detail='No URL supplied; no locator was added.'
   elif key in known_wrong:detail='The retrieved destination does not match the cited resource. '+str(ev.get('page_titles',[])[:2])
   elif status=='access_blocked':detail='Automated retrieval was blocked; page identity could not be confirmed. This is not evidence that the URL is broken.'
   elif status=='request_failed':detail='URL check could not complete: '+'; '.join(ev['notes'])
   else:detail='URL resolves, but the retrieved page did not establish a substantial match to the citation.'
   existing_missing=check=='missing_url' and any(r['citation_key']==key and r['check']=='missing_url' and r['status']=='open' for r in original)
   if not existing_missing:review.append([key,'accuracy',check,'url',detail,'open',None,ev.get('final_url') or ev['requested_url']])
   if c.value:font=copy(c.font);font.color='FFFF0000';c.font=font
   else:c.fill=PatternFill('solid',fgColor='FFFFC7CE')
   c.comment=Comment(detail,'Audit');counts['flagged_url_cells']+=1
  else:
   # Preserve previously unflagged links when this automated pass is inconclusive.
   counts['inconclusive_previously_unflagged']+=1
  if key in flagged:audit.append([key,'url_verification','flagged','matched' if matched else status,'Actual URL check replaces mutability-based screening',ev.get('final_url') or ev['requested_url']])
 uv=w.create_sheet('URL verification');uv.append(['citation_key','workbook_url','final_url','http_status','verification_status','record_match_signals','matched_authors','decision_notes'])
 for key,row in rows.items():
  ev=evidence[key];signals=', '.join(k for k,v in ev.get('record_matches',{}).items() if v)
  uv.append([key,ev['requested_url'],ev.get('final_url'),ev.get('http_status'),ev['status'],signals,'; '.join(ev.get('matched_author_names',[])),'; '.join(ev['notes'])])
 for c in uv[1]:c.font=Font(bold=True,color='FFFFFFFF');c.fill=PatternFill('solid',fgColor='243E50')
 uv.freeze_panes='A2';uv.auto_filter.ref=uv.dimensions
 for col,width in {'A':28,'B':90,'C':90,'D':15,'E':25,'F':45,'G':50,'H':100}.items():uv.column_dimensions[col].width=width
 policy='URL policy update: URLs are not flagged solely because their content is mutable, and an access date is not required to clear a URL flag. A resolving URL is accepted when the destination substantially identifies the cited work or resource through its title/core title, authors, identifiers, year, organization or project identity. Harmless title variations, updated repository titles and legitimate project homepages are accepted. Existing URL flags are cleared for verified matches; actual HTTP errors, mismatched destinations and unresolved existing URL checks remain flagged with specific reasons. Automated access blocks alone do not create new accuracy warnings on previously unflagged URLs. URL verification records all 637 live checks, including inconclusive results. URL strings and all other bibliography values and markings are preserved.'
 w['Policy'].append([policy])
 for a,b in zip(old['Bibliography'],s):
  for x,y in zip(a,b):
   assert x.value==y.value,x.coordinate
   if x.row>1 and x.column==ui+1:continue
   assert copy(x._style)==copy(y._style) and x.comment==y.comment,x.coordinate
 assert removed==132 and len(cleared)==84
 tmp=path.with_name('url-verification.tmp.xlsx');w.save(tmp);saved=openpyxl.load_workbook(tmp,rich_text=True)
 for r in saved['Bibliography'].iter_rows(min_row=2):
  ev=evidence[r[ki].value];c=r[ui]
  if ev['status']=='matched':assert c.font.color.rgb=='FF000000' and c.comment is None,r[ki].value
 tmp.replace(path)
 for sn,fn in [('Review','review.csv'),('Change audit','change-audit.csv'),('URL verification','url-verification.csv')]:
  with (B/fn).open('w',newline='',encoding='utf-8') as f:csv.writer(f).writerows(w[sn].values)
 with (B/'normalization-policy.txt').open('a') as f:f.write('\n'+policy+'\n')
 m=json.loads((B/'manifest.json').read_text());rv=[dict(zip(rh,r)) for r in list(review.values)[1:]]
 m['field_changes']=audit.max_row-1;m['review_status_counts']=dict(Counter(i['status'] for i in rv));m['review_counts']=dict(Counter(i['category'] for i in rv if i['status']=='open'));m['qc_counts']=dict(Counter(i['check'] for i in rv if i['status']=='open'));m['network_link_availability_checked']=True
 m['url_verification']={'records_checked':637,'status_counts':dict(Counter(ev['status'] for ev in evidence.values())),'mutability_only_findings_removed':removed,'previously_flagged_url_cells_cleared':len(cleared),**counts,'evidence_file':'url-verification-evidence.jsonl','url_values_unchanged_verified':True,'non_url_markings_unchanged_verified':True}
 (B/'manifest.json').write_text(json.dumps(m,indent=2)+'\n');print(json.dumps(m['url_verification'],indent=2))
if __name__=='__main__':main()
