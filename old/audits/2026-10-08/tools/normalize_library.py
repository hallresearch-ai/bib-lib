"""Create a separate normalized XLSX/BibTeX library, with auditable changes.

Requires openpyxl and bibtexparser >= 2. Run with --help for paths.
Existing source files and existing output directories are never overwritten.
"""

import argparse
import collections
import csv
import datetime as dt
import hashlib
import json
import re
import shutil
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

import bibtexparser
import openpyxl
from bibtexparser.middlewares.names import (
    parse_single_name_into_parts,
    split_multiple_persons_names,
)
from openpyxl.styles import Font, PatternFill


POLICY = """Normalization policy — 2026-10-08

The edited XLSX is the source of truth. The older BibTeX is used only to
recover provable spreadsheet type-conversion errors and to trace retired keys.
Source workbooks are copied byte-for-byte into snapshots; SHA-256 hashes are
recorded. No source file is overwritten. Empty spreadsheet rows are omitted.
Formula citation keys use their cached values and become literal text.

Citation keys remain stable. The key map includes every original and current
key, including previously deleted records. New keys, if needed in future,
should use FamilyYearShortTitle with deterministic alphabetical suffixes;
keys must not depend on row order. Previously deleted records are not restored.

Only identifier-and-version-confirmed duplicate records may be merged.
Same titles, proceedings ISBNs, homepages and related preprints are not sufficient.
The two Holistic Evaluation articles are retained: the three-author Annals
article and the multi-author TMLR article have different venues and scope.
The title-changing arXiv v1 prompt-injection preprint is retained separately
from the published conference citation. Two previously retired duplicate keys
are mapped to surviving keys; unrelated prior deletions are recorded explicitly.

DOIs are bare identifiers; ISBN punctuation is removed and original leading
zeros are recovered only from matching source BibTeX. PMID, PMCID, LCCN,
standard/report numbers stay in dedicated fields. arXiv identifiers belong in
eprint with archiveprefix=arXiv; explicit version suffixes are preserved.
URL normalization removes known tracking parameters, canonicalizes DOI and
arXiv locators, and uses HTTPS for known hosts. Arbitrary HTTP-only sites are
not assumed to support HTTPS. Record URLs are derived from existing locators;
approved URLs are accepted only for a unique title match and an empty URL.
An existing edited URL is not overwritten by a conflicting older proposal.
Duplicate locator-only notes/howpublished/bdsk-url-1 are removed after their
URL has been moved to url; substantive notes remain.

Types and field placement follow explicit metadata evidence. arXiv-only
articles become misc; journal Proceedings of the IEEE remains an article.
Conference publications use booktitle; reports use institution and number;
standards use standard with number. Unknown types remain review issues.
Journal names use full forms when the expansion is unambiguous. Case variants
and formatting-only wrappers are reconciled; ambiguous venue text is flagged.

Sentence case is used for titles. Acronyms, mixed-case technical names and a
declared proper-name vocabulary retain case inside braces. Ordinary words
lose cosmetic title-case braces; letter-by-letter acronym braces are joined.
Title wording, punctuation, math and TeX commands remain intact. Capitalization
changes are recorded and queued for proper-name review; this automated pass
does not certify official capitalization of every title.

Personal names use BibTeX's surname-first parser, preserving particles,
suffixes, accents and initials. No initials are expanded and no identities
are merged. Known institutional authors are braced as organizations, with
consistent recurring forms. Uncertain name/identity variants are review items.

Years contain four digits only, or are blank with an issue. Exact dates and
access dates use ISO YYYY-MM-DD; months use three-letter lowercase forms.
An Excel date in number/pages/note is recovered from the same original entry,
not interpreted as publication metadata. Missing years can be derived from an
existing explicit publication date. Dynamic years and year spans are retained
in the audit and cleared from year; no access year substitutes for publication.
Missing values stay blank. Placeholder publication metadata is removed and
logged. Numeric page ranges use --; article numbers remain article numbers.

QC distinguishes accuracy issues from cosmetic warnings. It checks required
fields by type; identifiers and checksums; duplicate identifiers/titles;
URLs; dates; name/venue variants; braces; formulas; export collisions; and
complete XLSX/BibTeX round-trip equality. Shared proceedings ISBNs are reported
as context, not duplicate proof. Web link availability and version fidelity are
not claimed: mutable/generic URLs and unverified replacements need review.
Nonstandard fields are preserved but flagged for downstream style support.
The Review sheet lists unresolved issues. This is a normalized working copy,
not a claim that all bibliographic facts have been externally verified.
"""

ORGANIZATIONS = {
    x.casefold(): x for x in [
        'National Institute of Standards and Technology', 'Microsoft',
        'Microsoft Research', 'H2O.ai', 'H2O.ai Team', 'OpenAI', 'IBM',
        'ASTM International', 'ISO', 'EU', 'FINRA', 'CFPB', 'FHFA', 'OCC',
        'ECB', 'OECD', 'OECD Artificial Intelligence Papers', 'Meta',
        'Frontier Model Forum', 'Hack the Future', 'HuggingFace', 'Adversa.ai',
        'CFRFM', 'JCGM', 'International Association of Privacy Professionals',
        'Scientific Integrity Framework Interagency Working Group of the National Science and Technology Council',
        'Gallup', 'Telescope Foundation',
    ]
}
ORGANIZATIONS['nist'] = 'National Institute of Standards and Technology'
VENUES = {
    'j. mach. learn. res.': 'Journal of Machine Learning Research',
    'the journal of machine learning research': 'Journal of Machine Learning Research',
    'trans. mach. learn. res.': 'Transactions on Machine Learning Research',
    'proc. acm hum.-comput. interact.': 'Proceedings of the ACM on Human-Computer Interaction',
    'j. data and information quality': 'Journal of Data and Information Quality',
    'res involv engagem': 'Research Involvement and Engagement',
    'j bus ethics': 'Journal of Business Ethics',
    'acm computing surveys (csur)': 'ACM Computing Surveys',
    'international journal of computer vision (ijcv)': 'International Journal of Computer Vision',
    'acm transactions on graphics (tog)': 'ACM Transactions on Graphics',
}
PROPER = set('''ChatGPT RoBERTa Microsoft H2O H2O.ai Python GitHub JIRA NIST
Google Facebook Amazon Anthropic OpenAI IBM NVIDIA Intel Stanford Harvard
ARIA TabArena LiveBench Promptbench PromptBench Taskbench TaskBench AudioCaps
ImageNet MNIST BRATS BLEU GPT GPT-4 GPT-3 AI NLP ML LLM LLMs GenAI XAI
Wikipedia YouTube Bing Feynman Heilmeier Molnar Gallup Telescope NeurIPS
Wunsch-Bell Dwyer Maloney Tukey Bayes Bayesian Markov Shapley Banzhaf
Gini Fisher Pearson Spearman Kendall Krippendorff Cronbach Likert Holm
Bonferroni Kolmogorov Smirnov Monte Carlo Gaussian Fourier LeCun Cortes
Salakhutdinov Hinton Hawkes Poisson Bernoulli Euler Laplace Lagrange
SAS HMDA R CUDA Java RStudio Fairlearn FairML SHAP LIME AdaBoost XGBoost
LightGBM CatBoost AutoML AutoGluon PyTorch TensorFlow pandas NumPy sklearn
scikit-learn Kubernetes Linux BSD SPDX Presidio ConductorAI DeepSeek
Llama BERT Turing American America Americans Europe European United
States Norway Norwegian Britain British English China Chinese India
Indian UK U.S. US ACM IEEE ISO IEC ASTM CFPB OECD NIH NHS GDPR HAX
Rorschach Hermeneutic Prometheus Dalmatian KDD'''.split())
KNOWN_HTTPS = {'arxiv.org', 'www.arxiv.org', 'doi.org', 'dx.doi.org',
                'h2o.ai', 'www.h2o.ai', 'docs.h2o.ai', 'data.h2o.ai',
                'www.jmlr.org', 'jmlr.org', 'github.com'}
MONTHS = {name.lower(): f'{n:02d}' for n, name in enumerate(
    ['January', 'February', 'March', 'April', 'May', 'June', 'July',
     'August', 'September', 'October', 'November', 'December'], 1)}
MONTH_ABBR = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']
DATE_FIELDS = {'date', 'urldate', 'issue_date'}


def text(value):
    if value is None:
        return ''
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def title_id(value):
    return re.sub(r'[^\w]', '', value.casefold().replace('_', ''))


def clean_url(value):
    value = value.strip()
    wrapped = re.fullmatch(r'\\+url\{(.*)\}', value)
    if wrapped:
        value = wrapped[1]
    try:
        p = urlsplit(value)
        host = p.hostname or ''
        scheme = 'https' if host in KNOWN_HTTPS else p.scheme
        query = [(k, v) for k, v in parse_qsl(p.query, keep_blank_values=True)
                 if not k.lower().startswith('utm_') and k.lower() not in
                 {'fbclid', 'gclid', 'msclkid', 'mc_cid', 'mc_eid', 'trk'}]
        path = p.path
        netloc = p.netloc
        if host in {'doi.org', 'dx.doi.org'}:
            netloc = 'doi.org'
        if host in {'arxiv.org', 'www.arxiv.org'}:
            netloc = 'arxiv.org'
            path = re.sub(r'^/pdf/', '/abs/', path)
            path = re.sub(r'\.pdf$', '', path)
        return urlunsplit((scheme, netloc, path, urlencode(query), p.fragment))
    except ValueError:
        return value


def doi_id(value):
    return re.sub(r'^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)', '', value, flags=re.I).strip()


def sentence_title(value):
    # Whole-title braces are formatting, not a proper-name assertion.
    if value.startswith('{') and value.endswith('}'):
        depth = 0
        outer = True
        for i, c in enumerate(value):
            if c == '{': depth += 1
            if c == '}': depth -= 1
            if depth == 0 and i < len(value) - 1: outer = False
        if outer: value = value[1:-1]
    # Join letter-by-letter acronym protection before normalizing words.
    value = re.sub(r'(?:\{[A-Z]\}){2,}', lambda m: '{' + re.sub(r'[{}]', '', m[0]) + '}', value)
    # Unwrap single-letter title-case decoration, e.g. {M}achine.
    value = re.sub(r'\{([A-Za-z])\}(?=[a-z])', r'\1', value)
    pattern = r'\$[^$]*\$|\\[A-Za-z]+(?:\{[^{}]*\})?|\{[^{}]*\}|[\w]+(?:[.\-/][\w]+)*'
    first = True
    def replace(m):
        nonlocal first
        word = m[0]
        if word.startswith(('\\', '$')):
            first = False
            return word
        braced = word.startswith('{')
        inner = word[1:-1] if braced else word
        is_first = first
        first = False
        acronyms = {'ai': 'AI', 'ml': 'ML', 'nlp': 'NLP', 'llm': 'LLM', 'llms': 'LLMs', 'h2o': 'H2O'}
        if inner.lower() in acronyms: inner = acronyms[inner.lower()]
        preserve = inner in PROPER or (
            any(c.isalpha() for c in inner) and inner.isupper() and len(inner) > 1
        ) or bool(re.search(r'(?<=[a-z])[A-Z]', inner))
        if preserve:
            return '{' + inner + '}'
        result = inner.lower()
        if is_first and result: result = result[0].upper() + result[1:]
        return result
    return re.sub(pattern, replace, value)


def braces_balanced(value):
    depth = 0
    for m in re.finditer(r'(?<!\\)[{}]', value):
        depth += 1 if m[0] == '{' else -1
        if depth < 0: return False
    return depth == 0


def isbn_valid(value):
    if re.fullmatch(r'\d{9}[\dX]', value):
        return sum((10-i)*(10 if c == 'X' else int(c)) for i,c in enumerate(value)) % 11 == 0
    if re.fullmatch(r'\d{13}', value):
        return sum((1 if i % 2 == 0 else 3)*int(c) for i,c in enumerate(value)) % 10 == 0
    return False


def write_csv(path, columns, rows):
    with path.open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def run(source, old_bib, approval, output):
    if output.exists():
        raise ValueError(f'Refusing to overwrite existing output directory: {output}')
    ws = openpyxl.load_workbook(source, data_only=True).active
    values = list(ws.values)
    source_row_count = len(values) - 1
    headers = list(values[0])
    assert len(headers) == len(set(headers)) and None not in headers
    rows = [dict(zip(headers, r)) for r in values[1:] if any(v is not None for v in r)]
    parsed_original = bibtexparser.parse_file(str(old_bib))
    if parsed_original.failed_blocks:
        raise ValueError('Original BibTeX has parser failures')
    original = {e.key: {f.key: f.value for f in e.fields} for e in parsed_original.entries}
    current_keys = [r['citation_key'] for r in rows]
    assert all(current_keys) and len(current_keys) == len(set(current_keys))
    audit, issues = [], []
    for row in openpyxl.load_workbook(source).active:
        for cell in row:
            if cell.data_type == 'f':
                key = values[cell.row-1][headers.index('citation_key')]
                audit.append(dict(citation_key=key, field=headers[cell.column-1],
                                  before=cell.value, after=text(values[cell.row-1][cell.column-1]),
                                  reason='Replace spreadsheet formula with cached literal value', evidence=str(source)))
    def issue(key, category, check, field, detail):
        issues.append(dict(citation_key=key, category=category, check=check, field=field, detail=detail))
    def change(row, field, value, reason, evidence='source workbook'):
        before = text(row.get(field))
        after = text(value)
        if before != after:
            audit.append(dict(citation_key=row['citation_key'], field=field, before=before,
                              after=after, reason=reason, evidence=evidence))
        row[field] = after
        if field not in headers: headers.append(field)

    proposals = collections.defaultdict(list)
    for title, url, decision in list(openpyxl.load_workbook(approval, data_only=True).active.values)[1:]:
        proposals[title_id(text(title))].append((text(url), text(decision)))
    title_counts = collections.Counter(title_id(text(r['title'])) for r in rows)
    # Canonical forms chosen for formatting-only case variants.
    venues = collections.defaultdict(collections.Counter)
    for r in rows:
        value = re.sub(r'\\+textit\{(.*?)\}', r'\1', text(r.get('journal')))
        if value: venues[value.casefold()][value] += 1
    for group, choices in venues.items():
        VENUES.setdefault(group, max(choices, key=lambda v: (sum(c.isupper() for c in v), choices[v], v)))

    for r in rows:
        key = r['citation_key']
        for field, raw in list(r.items()):
            if isinstance(raw, (dt.date, dt.datetime)):
                if field in DATE_FIELDS:
                    change(r, field, raw.date().isoformat() if isinstance(raw, dt.datetime) else raw.isoformat(), 'Convert Excel date to ISO text')
                elif field == 'year':
                    change(r, field, str(raw.year), 'Extract four-digit publication year')
                    if not r.get('month'): change(r, 'month', MONTH_ABBR[raw.month-1], 'Extract month from existing publication date')
                elif field == 'month':
                    change(r, 'month', MONTH_ABBR[raw.month-1], 'Extract month from existing publication timestamp')
                    if not r.get('date'): change(r, 'date', raw.date().isoformat(), 'Preserve publication date from month timestamp')
                else:
                    prior = original.get(key, {}).get(field)
                    if prior is not None:
                        change(r, field, prior, 'Recover Excel-converted text from matching original entry', str(old_bib))
                    else:
                        issue(key, 'accuracy', 'spreadsheet_conversion', field, text(raw))
            # All cells in normalized bibliography are literal strings.
            r[field] = text(r[field])

        for field in ('author', 'editor'):
            value = re.sub(r'\s+', ' ', r.get(field, '')).strip()
            if not value: continue
            parts = []
            whole_org = ORGANIZATIONS.get(value.strip('{} ').casefold())
            authors = ['{' + whole_org + '}'] if whole_org else split_multiple_persons_names(value)
            for name in authors:
                bare = name.strip('{} ')
                if bare.casefold() in ORGANIZATIONS:
                    parts.append('{' + ORGANIZATIONS[bare.casefold()] + '}')
                elif name.startswith('{') and name.endswith('}') or bare in {'others', 'others.'}:
                    parts.append('others' if bare.startswith('others') else name)
                else:
                    try: parts.append(parse_single_name_into_parts(name).merge_last_name_first)
                    except (ValueError, IndexError):
                        parts.append(name)
                        issue(key, 'accuracy', 'author_parse', field, name)
            change(r, field, ' and '.join(parts), 'Use surname-first names and protect known institutional authors')

        doi = doi_id(r.get('doi', ''))
        if doi: change(r, 'doi', doi, 'Store bare DOI')
        isbn = re.sub(r'[-\s]', '', r.get('isbn', '')).upper()
        if isbn and not isbn_valid(isbn):
            candidate = re.sub(r'[-\s]', '', original.get(key, {}).get('isbn', '')).upper()
            if isbn_valid(candidate) and isbn.lstrip('0') == candidate.lstrip('0'):
                isbn = candidate
        if isbn: change(r, 'isbn', isbn, 'Store ISBN without punctuation; recover provable lost leading zeros', str(old_bib))

        url = r.get('url', '')
        approved = [u for u, decision in proposals.get(title_id(r.get('title', '')), []) if decision == 'Y' and u]
        if not url and len(approved) == 1 and title_counts[title_id(r['title'])] == 1:
            url = approved[0]
        if not url:
            for field in ('bdsk-url-1', 'howpublished', 'note', 'journal'):
                found = re.search(r'https?://[^\s{}]+', r.get(field, ''))
                if found:
                    url = found[0]
                    break
        if not url and doi: url = 'https://doi.org/' + doi
        if url: change(r, 'url', clean_url(url), 'Normalize or extract existing/approved record locator')
        if not r.get('doi') and re.match(r'https://doi\.org/10\.', r.get('url', '')):
            change(r, 'doi', r['url'].split('doi.org/', 1)[1], 'Extract DOI from existing DOI landing URL')
        for field in ('bdsk-url-1', 'howpublished', 'note'):
            value = r.get(field, '')
            match = re.fullmatch(r'(?:URL:\s*)?(?:\\+url\{)?(https?://[^{}\s]+)\}?', value, re.I)
            if match and clean_url(match[1]) == r.get('url'):
                change(r, field, '', 'Remove duplicated locator-only field after retaining url')

        arxiv = None
        for field in ('url', 'doi', 'journal', 'howpublished', 'eprint'):
            v = r.get(field, '')
            m = re.search(r'(?:arxiv(?:\.org/(?:abs|pdf)/|[.:/ ]+)|^)(\d{4}\.\d{4,5}(?:v\d+)?|[a-z-]+/\d{7}(?:v\d+)?)', v, re.I)
            if m:
                arxiv = m[1]
                break
        if arxiv:
            change(r, 'eprint', arxiv, 'Use dedicated arXiv identifier; preserve explicit version')
            change(r, 'archiveprefix', 'arXiv', 'Use canonical archive name')
            if r.get('journal', '').lower().startswith(('arxiv', 'corr')):
                change(r, 'journal', '', 'Move arXiv venue into archive metadata')
                if r['entry_type'] == 'article': change(r, 'entry_type', 'misc', 'Classify arXiv-only article as preprint')
            if r['entry_type'] == 'article' and not r.get('journal') and not r.get('booktitle'):
                change(r, 'entry_type', 'misc', 'Classify record with only arXiv publication evidence as preprint')
        elif r.get('eprint', '').startswith(('http://', 'https://')):
            change(r, 'pdf', r['eprint'], 'Move PDF URL out of identifier field')
            change(r, 'eprint', '', 'Retain full-text PDF locator in pdf')

        year = r.get('year', '')
        if year and not re.fullmatch(r'\d{4}', year):
            issue(key, 'accuracy', 'unverified_year', 'year', f'Original value: {year}; left blank, not guessed')
            change(r, 'year', '', 'Remove non-year placeholder/span; retain original in audit')
        date = r.get('date', '')
        if not r.get('year') and re.fullmatch(r'\d{4}(?:-\d{2}(?:-\d{2})?)?', date):
            change(r, 'year', date[:4], 'Derive year from explicit publication date')
        month = r.get('month', '').lower().rstrip('.')
        month_num = MONTHS.get(month)
        if not month_num:
            for n, abbr in enumerate(MONTH_ABBR, 1):
                if month == abbr or month == 'sept' and n == 9: month_num = f'{n:02d}'
        if month.isdigit() and 1 <= int(month) <= 12: month_num = f'{int(month):02d}'
        md = re.fullmatch(r'([A-Za-z]+)~?(\d{1,2})', month)
        if md and md[1] in MONTHS:
            month_num = MONTHS[md[1]]
            if r.get('year'):
                change(r, 'date', f"{r['year']}-{month_num}-{int(md[2]):02d}", 'Preserve explicit month/day as publication date')
        if month_num: change(r, 'month', MONTH_ABBR[int(month_num)-1], 'Use three-letter month')
        if not r.get('urldate'):
            m = re.search(r'accessed\s*[:=]?\s*(\d{4}-\d{2}-\d{2})', r.get('note', ''), re.I)
            if not m:
                m2 = re.search(r'accessed\s*[:=]?\s*([A-Za-z]+)\s+(\d{1,2}),?\s+(\d{4})', r.get('note', ''), re.I)
                if m2 and m2[1].lower() in MONTHS:
                    change(r, 'urldate', f'{m2[3]}-{MONTHS[m2[1].lower()]}-{int(m2[2]):02d}', 'Extract explicitly recorded access date')
            else: change(r, 'urldate', m[1], 'Extract explicitly recorded access date')

        # Verified type/field repairs; links are included in the audit.
        if key == 'GradientBasedLearnin278':
            evidence = 'https://bottou.org/papers/lecun-98h'
            for field, val in dict(entry_type='article', journal='Proceedings of the IEEE', booktitle='', volume='86', number='11', pages='2278--2324', year='1998', month='nov').items():
                change(r, field, val, 'Match author-hosted journal publication record', evidence)
        if r['entry_type'] == 'paper' and r.get('institution'):
            change(r, 'entry_type', 'techreport', 'Institutional repository classifies item as report', 'https://doi.org/10.26188/28822919')
        if key == 'ArtificialIntelligen078':
            evidence = 'https://doi.org/10.6028/NIST.AI.600-1'
            for field, val in dict(entry_type='techreport', institution='National Institute of Standards and Technology', number='AI 600-1', publisher='').items():
                change(r, field, val, 'Use NIST report fields', evidence)
        if key == 'ArtificialIntelligen079':
            change(r, 'entry_type', 'book', 'Classify ISBN-bearing textbook as book', 'source ISBN and publisher')
        if key == 'WhichHumans635':
            change(r, 'entry_type', 'misc', 'Author publication page identifies working paper', 'https://henrich.fas.harvard.edu/publication/which-humans')
            change(r, 'year', '2023', 'Use publication year on author page', 'https://henrich.fas.harvard.edu/publication/which-humans')
        journal = re.sub(r'\\+textit\{(.*?)\}', r'\1', r.get('journal', ''))
        if journal:
            change(r, 'journal', VENUES.get(journal.casefold(), journal), 'Use full canonical journal name; remove formatting wrapper')
        for field in ('booktitle', 'publisher', 'institution'):
            v = re.sub(r'\\+textit\{(.*?)\}', r'\1', r.get(field, ''))
            if v == 'NIST': v = 'National Institute of Standards and Technology'
            if v: change(r, field, v, 'Normalize publication metadata formatting')
        if r.get('publisher') == 'Publisher name, location':
            issue(key, 'accuracy', 'placeholder', 'publisher', r['publisher'])
            change(r, 'publisher', '', 'Remove literal publisher placeholder')
        pages = re.sub(r'(?<=\d)\s*[-–—]\s*(?=\d)', '--', r.get('pages', ''))
        if pages: change(r, 'pages', pages, 'Use BibTeX numeric page ranges')
        if r.get('issue') and not r.get('number'):
            change(r, 'number', r['issue'], 'Move issue to BibTeX number')
            change(r, 'issue', '', 'Remove moved issue field')
        before = r.get('title', '')
        after = sentence_title(before)
        if after != before:
            change(r, 'title', after, 'Apply sentence case with protected terms; preserve wording')
            issue(key, 'cosmetic', 'title_capitalization_review', 'title', 'Confirm proper-name protection after automated sentence case')

    # Quality-control passes are independent of mutation decisions.
    required = {
        'article': [('title',), ('author',), ('journal',), ('year',)],
        'inproceedings': [('title',), ('author',), ('booktitle',), ('year',)],
        'incollection': [('title',), ('author',), ('booktitle',), ('publisher',), ('year',)],
        'book': [('title',), ('author', 'editor'), ('publisher',), ('year',)],
        'techreport': [('title',), ('author',), ('institution',), ('year',)],
        'manual': [('title',)], 'standard': [('title',), ('number',)],
        'online': [('title',), ('url',)], 'misc': [('title',)],
        'proceedings': [('title',), ('year',)],
    }
    ids = {field: collections.defaultdict(list) for field in ('doi', 'isbn', 'pmid', 'pmcid', 'eprint', 'title')}
    names = collections.defaultdict(set)
    for r in rows:
        key = r['citation_key']
        if r['entry_type'] not in required: issue(key, 'accuracy', 'unknown_type', 'entry_type', r['entry_type'])
        for alternatives in required.get(r['entry_type'], []):
            if not any(r.get(f) for f in alternatives): issue(key, 'accuracy', 'missing_required', '|'.join(alternatives), r['entry_type'])
        for field, groups in ids.items():
            v = r.get(field, '')
            if v: groups[title_id(v) if field == 'title' else v.casefold()].append(key)
        for field, value in r.items():
            if not braces_balanced(value): issue(key, 'accuracy', 'unbalanced_braces', field, value)
            if value.startswith('='): issue(key, 'accuracy', 'formula_in_export', field, value)
        doi = r.get('doi', '')
        if doi and not re.fullmatch(r'10\.\d{4,9}/[^\s{}]+', doi): issue(key, 'accuracy', 'doi_syntax', 'doi', doi)
        isbn = r.get('isbn', '')
        if isbn and not isbn_valid(isbn): issue(key, 'accuracy', 'isbn_checksum', 'isbn', isbn)
        for field, pattern in [('pmid', r'\d+'), ('pmcid', r'PMC\d+'), ('eprint', r'(?:\d{4}\.(?:\d{4,5})|[a-z-]+/\d{7})(?:v\d+)?')]:
            if r.get(field) and not re.fullmatch(pattern, r[field], re.I): issue(key, 'accuracy', 'identifier_syntax', field, r[field])
        if r.get('issn'):
            for raw_issn in r['issn'].split(','):
                issn = re.sub('[- ]', '', raw_issn).upper()
                if not re.fullmatch(r'\d{7}[\dX]', issn) or sum((8-i)*(10 if c == 'X' else int(c)) for i,c in enumerate(issn)) % 11:
                    issue(key, 'accuracy', 'issn_checksum', 'issn', raw_issn)
            if ',' in r['issn']: issue(key, 'cosmetic', 'multiple_issns', 'issn', 'Valid multiple ISSNs preserved; distinguish print/electronic after verification')
        for field in DATE_FIELDS:
            v = r.get(field, '')
            if v and not re.fullmatch(r'\d{4}(?:-\d{2}(?:-\d{2})?)?', v): issue(key, 'accuracy', 'date_syntax', field, v)
            elif len(v) == 10:
                try: dt.date.fromisoformat(v)
                except ValueError: issue(key, 'accuracy', 'invalid_date', field, v)
        if r.get('year') and not 1500 <= int(r['year']) <= dt.date.today().year + 1:
            issue(key, 'accuracy', 'suspicious_year', 'year', r['year'])
        if r.get('month') and r['month'] not in MONTH_ABBR: issue(key, 'accuracy', 'month_syntax', 'month', r['month'])
        url = r.get('url', '')
        if not url: issue(key, 'accuracy', 'missing_url', 'url', 'No record locator available')
        else:
            try:
                p = urlsplit(url)
                if p.scheme not in {'http', 'https'} or not p.hostname or re.search(r'[\s{}<>]', url):
                    issue(key, 'accuracy', 'malformed_url', 'url', url)
                if p.path in {'', '/', '/resources'}: issue(key, 'accuracy', 'generic_landing_page', 'url', url)
                if r['entry_type'] in {'online', 'manual', 'misc'} and not r.get('eprint') and not r.get('doi'):
                    issue(key, 'accuracy', 'mutable_web_resource', 'url', 'Confirm cited version and access date; live availability not checked')
            except ValueError: issue(key, 'accuracy', 'malformed_url', 'url', url)
        for name in split_multiple_persons_names(r.get('author', '')):
            if ',' in name and not name.startswith('{'):
                family, given = name.split(',', 1)
                initials = ''.join(p[0].lower() for p in given.split() if p)
                names[(family.casefold(), initials)].add((name, key))
        journal = r.get('journal', '')
        if journal and ('http' in journal or 'arxiv' in journal.casefold() or 'Proceedings of SAS' in journal or 'Advances in Neural' in journal):
            issue(key, 'accuracy', 'venue_field_review', 'journal', journal)
        for field in ('file', 'abstractnote', 'author1_email', 'author2_email', 'bdsk-url-1', 'issue_date'):
            if r.get(field): issue(key, 'cosmetic', 'nonstandard_field', field, 'Preserved; verify downstream bibliography style supports this field')
    for field, groups in ids.items():
        for value, keys in groups.items():
            if len(keys) > 1:
                category = 'cosmetic' if field == 'isbn' else 'accuracy'
                detail = ', '.join(keys)
                if field == 'isbn': detail += '; may identify shared proceedings, not duplicate papers'
                if set(keys) == {'HolisticEvaluationof301', 'HolisticEvaluationof302'}:
                    category = 'context'
                    detail += '; retained distinct Annals/TMLR publications: https://doi.org/10.1111/nyas.15007 ; https://mlanthology.org/tmlr/2023/liang2023tmlr-holistic/'
                for key in keys: issue(key, category, 'shared_' + field, field, detail)
    for values in names.values():
        if len({n for n,k in values}) > 1:
            for name,key in sorted(values): issue(key, 'accuracy', 'author_variant', 'author', '; '.join(sorted({n for n,k in values})))

    key_map = []
    retired = {'Notwhatyouvesignedup412': 'Notwhatyouvesignedup411',
               'TheAssessingRisksand558': 'TheAssessingRisksand557'}
    for key in sorted(set(original) | set(current_keys)):
        new = key if key in current_keys else retired.get(key, '')
        status = 'preserved' if key in current_keys else 'merged_in_source' if new else 'deleted_in_source'
        key_map.append(dict(old_key=key, new_key=new, status=status))
        if key not in current_keys:
            audit.append(dict(citation_key=key, field='citation_key', before=key, after=new,
                              reason=status + '; prior workbook decision preserved', evidence=str(source)))

    # Fail before writing when a structural defect would corrupt an export.
    fatal = [i for i in issues if i['check'] in {'unbalanced_braces', 'formula_in_export'}]
    if fatal: raise ValueError(f'Unsafe export fields: {fatal}')
    control = {'entry_type', 'citation_key'}
    bib = '\n\n'.join('@' + r['entry_type'] + '{' + r['citation_key'] + ',\n' +
            ',\n'.join('  ' + f + ' = {' + r.get(f, '') + '}' for f in headers
                       if f not in control and r.get(f)) + '\n}' for r in rows) + '\n'
    parsed = bibtexparser.parse_string(bib)
    assert not parsed.failed_blocks, parsed.failed_blocks
    assert len(parsed.entries) == len(rows)
    roundtrip = {e.key: (e.entry_type, {f.key: f.value for f in e.fields}) for e in parsed.entries}
    for r in rows:
        assert roundtrip[r['citation_key']] == (r['entry_type'], {f:r[f] for f in headers if f not in control and r.get(f)}), r['citation_key']

    output.mkdir(parents=True)
    snapshots = output / 'snapshots'
    snapshots.mkdir()
    source_hashes = {}
    for p in (source, old_bib, approval):
        shutil.copy2(p, snapshots / p.name)
        source_hashes[str(p)] = hashlib.sha256(p.read_bytes()).hexdigest()
        assert hashlib.sha256((snapshots / p.name).read_bytes()).hexdigest() == source_hashes[str(p)]
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    def sheet(name, columns, records):
        ws = workbook.create_sheet(name)
        ws.append(columns)
        for r in records:
            ws.append([r.get(c, '') for c in columns])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='243E50')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(70, max(15, max(len(text(c.value)) for c in col[:30]) + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.number_format = '@'
                if c.value is not None: c.data_type = 's'
        return ws
    sheet('Bibliography', headers, rows)
    sheet('Change audit', ['citation_key','field','before','after','reason','evidence'], audit)
    sheet('Review', ['citation_key','category','check','field','detail'], issues)
    sheet('Citation key map', ['old_key','new_key','status'], key_map)
    policy_rows = [{'policy':line} for line in POLICY.splitlines()]
    sheet('Policy', ['policy'], policy_rows)
    workbook.save(output / 'library-cleaned.xlsx')
    (output / 'library-cleaned.bib').write_text(bib, encoding='utf-8')
    (output / 'normalization-policy.txt').write_text(POLICY, encoding='utf-8')
    write_csv(output / 'change-audit.csv', ['citation_key','field','before','after','reason','evidence'], audit)
    write_csv(output / 'review.csv', ['citation_key','category','check','field','detail'], issues)
    write_csv(output / 'citation-key-map.csv', ['old_key','new_key','status'], key_map)
    saved = list(openpyxl.load_workbook(output / 'library-cleaned.xlsx', data_only=True)['Bibliography'].values)
    assert list(saved[0]) == headers
    for expected, actual in zip(rows, saved[1:], strict=True):
        assert [expected.get(f, '') for f in headers] == [text(v) for v in actual]
    manifest = dict(source_hashes=source_hashes, records=len(rows), blank_rows_omitted=source_row_count-len(rows),
                    field_changes=len(audit), review_counts=dict(collections.Counter(i['category'] for i in issues)),
                    qc_counts=dict(collections.Counter(i['check'] for i in issues)),
                    roundtrip_verified=True, source_snapshots_verified=True,
                    network_link_availability_checked=False)
    (output / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(manifest, indent=2))
    print(f'Created {output / "library-cleaned.xlsx"}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=Path('bib_files/lib-2026-07-17_1632_codex.xlsx'))
    parser.add_argument('--original-bib', type=Path, default=Path('bib_files/lib-2026-07-17_1632.bib'))
    parser.add_argument('--approvals', type=Path, default=Path('bib_files/bib-url-approval-2026-10-08.xlsx'))
    parser.add_argument('--output', type=Path, default=Path('bib_files/normalized-2026-10-08'))
    args = parser.parse_args()
    run(args.source, args.original_bib, args.approvals, args.output)
